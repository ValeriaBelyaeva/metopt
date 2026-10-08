"""Excel для ЛР2: ручное прохождение метода ломаных (каждая итерация — строка с формулами)."""
import sys
from pathlib import Path

import openpyxl
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

LAB = Path(__file__).resolve().parent.parent / "lr2"
sys.path.insert(0, str(LAB))
from python_solution import parse_function, piyavskii  # noqa: E402

A, B, L, EPS = -3, 3, 1 + 3.14159, 0.01
res = piyavskii(parse_function("x + sin(3.14159*x)"), A, B, L, EPS)
hist = res["history"]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Метод ломаных"
thin = Side(style="thin", color="999999")
BRD = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="DCE6F1")
PAR = PatternFill("solid", fgColor="FFF2CC")
STOP = PatternFill("solid", fgColor="E2EFDA")
bold = Font(bold=True)

ws["A1"] = "ЛР2. Метод ломаных — ручное прохождение: f(x) = x + sin(3.14159·x) на [-3; 3]"
ws["A1"].font = Font(bold=True, size=13)
params = [("a", A), ("b", B), ("L", "=1+3.14159"), ("eps", EPS),
          ("f(a)", "=B3+SIN(3.14159*B3)"), ("f(b)", "=B4+SIN(3.14159*B4)")]
for i, (k, v) in enumerate(params):
    ws.cell(row=3 + i, column=1, value=k).font = bold
    c = ws.cell(row=3 + i, column=2, value=v)
    c.fill = PAR
    c.border = BRD
notes = [
    "L = max|f'(x)| = max|1 + π·cos(πx)| = 1 + π.",
    "Нижняя вершина ломаной на [xл; xп]:  x̂ = (xл + xп)/2 + (f(xл) − f(xп))/(2L),   p̂ = (f(xл) + f(xп))/2 − L(xп − xл)/2.",
    "На каждом шаге берём самую низкую из ещё не использованных вершин (её интервал вписан в строку),",
    "вычисляем f(x̂) — это новая точка испытания; её вершина заменяется двумя новыми (слева и справа).",
    "Остановка: рекорд f* − p̂ ≤ eps (p̂ — минимум ломаной, f* — лучшее найденное значение).",
]
for i, t in enumerate(notes):
    ws.cell(row=3 + i, column=4, value=t).font = Font(italic=True, color="555555")

HR = 10
head = ["k", "точки (№)", "xл", "f(xл)", "xп", "f(xп)", "x̂ — вершина", "p̂ — высота", "рекорд f*",
        "f* − p̂", "f(x̂)", "№ новой точки", "статус"]
for j, t in enumerate(head, 1):
    c = ws.cell(row=HR, column=j, value=t)
    c.font = bold
    c.fill = HEAD
    c.border = BRD
    c.alignment = Alignment(horizontal="center", wrap_text=True)


def ref(i):
    """Ячейки (x, f) точки испытания с номером i: 0 = a, 1 = b, k+1 — точка k-й итерации."""
    if i == 0:
        return "$B$3", "$B$7"
    if i == 1:
        return "$B$4", "$B$8"
    r = HR + (i - 1)
    return f"$G${r}", f"$K${r}"


for h in hist:
    k = h["k"]
    r = HR + k
    xl, fl = ref(h["il"])
    xr, fr = ref(h["ir"])
    row = [k, f"{h['il']} и {h['ir']}", f"={xl}", f"={fl}", f"={xr}", f"={fr}",
           f"=(C{r}+E{r})/2+(D{r}-F{r})/(2*$B$5)",
           f"=(D{r}+F{r})/2-$B$5*(E{r}-C{r})/2",
           f"=MIN($B$7,$B$8,K${HR}:K{r - 1})",
           f"=I{r}-H{r}",
           None if h["stop"] else f"=G{r}+SIN(3.14159*G{r})",
           None if h["stop"] else k + 1,
           f'=IF(J{r}<=$B$6,"СТОП: f*-p̂ <= eps","испытание")']
    for j, v in enumerate(row, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.border = BRD
        if j in (3, 4, 5, 6, 7, 8, 9, 10, 11):
            c.number_format = "0.00000"
        if h["stop"]:
            c.fill = STOP

last = HR + len(hist)
ans = last + 2
ws.cell(row=ans, column=1, value="Ответ").font = Font(bold=True, size=12)
ws.cell(row=ans + 1, column=1, value="f(x*) ≈")
ws.cell(row=ans + 1, column=2, value=f"=I{last}").number_format = "0.00000"
ws.cell(row=ans + 2, column=1, value="x* ≈")
ws.cell(row=ans + 2, column=2,
        value=f"=INDEX(G{HR + 1}:G{last - 1},MATCH(I{last},K{HR + 1}:K{last - 1},0))").number_format = "0.00000"
ws.cell(row=ans + 3, column=1, value="итераций")
ws.cell(row=ans + 3, column=2, value=f"=COUNT(K{HR + 1}:K{last})")
ws.cell(row=ans + 4, column=1, value="точность f* − p̂")
ws.cell(row=ans + 4, column=2, value=f"=J{last}").number_format = "0.00000"
for i in range(4):
    ws.cell(row=ans + 1 + i, column=2).fill = STOP
    ws.cell(row=ans + 1 + i, column=1).font = bold

widths = [6, 10, 10, 10, 10, 10, 13, 12, 11, 10, 10, 9, 22]
for j, w in enumerate(widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
ws.freeze_panes = ws.cell(row=HR + 1, column=1)

# ------------------------------------------------ лист с графиком
g = wb.create_sheet("График")
g["A1"], g["B1"] = "x", "f(x)"
N = 240
for i in range(N + 1):
    g.cell(row=2 + i, column=1, value=f"={A}+{(B - A) / N}*{i}")
    g.cell(row=2 + i, column=2, value=f"=A{2 + i}+SIN(3.14159*A{2 + i})")
ch = ScatterChart()
ch.title = "f(x) и точки испытаний метода ломаных"
ch.x_axis.title = "x"
ch.y_axis.title = "f(x)"
ch.height, ch.width = 10, 20
s1 = Series(Reference(g, min_col=2, min_row=2, max_row=N + 2),
            Reference(g, min_col=1, min_row=2, max_row=N + 2), title="f(x)")
s1.marker.symbol = "none"
s1.smooth = True
s2 = Series(Reference(ws, min_col=11, min_row=HR + 1, max_row=last - 1),
            Reference(ws, min_col=7, min_row=HR + 1, max_row=last - 1), title="точки испытаний")
s2.marker.symbol = "circle"
s2.marker.size = 5
s2.graphicalProperties.line.noFill = True
ch.series += [s1, s2]
g.add_chart(ch, "D2")

wb.save(LAB / "excel_solution.xlsx")
print("saved", len(hist), "rows; iterations =", res["iterations"])
