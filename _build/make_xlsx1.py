import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName

wb = openpyxl.Workbook()
thin = Side(style="thin", color="999999"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="DCE6F1"); VAR = PatternFill("solid", fgColor="FFF2CC"); OBJ = PatternFill("solid", fgColor="E2EFDA")
bold = Font(bold=True)

def solver_model(ws, idx, opt, adj, typ, cons, neg=1):
    """Скрытые имена, в которых Excel хранит модель «Поиска решения» листа.
    rel: 1 '<=', 2 '=', 3 '>=', 4 'int'."""
    q = f"'{ws.title}'!"
    names = {
        "solver_opt": q + opt, "solver_adj": q + adj, "solver_typ": str(typ), "solver_val": "0",
        "solver_eng": "2", "solver_lin": "1", "solver_neg": str(neg), "solver_num": str(len(cons)),
        "solver_ver": "3", "solver_cvg": "0.0001", "solver_drv": "1", "solver_est": "1",
        "solver_itr": "2147483647", "solver_mip": "2147483647", "solver_mni": "30", "solver_mrt": "0.075",
        "solver_msl": "2", "solver_nod": "2147483647", "solver_pre": "0.000001", "solver_rbv": "1",
        "solver_rlx": "2", "solver_rsd": "0", "solver_scl": "1", "solver_sho": "2", "solver_ssz": "100",
        "solver_tim": "2147483647", "solver_tol": "0", "solver_nwt": "1",
    }
    for k, (lhs, rel, rhs) in enumerate(cons, 1):
        names[f"solver_lhs{k}"] = q + lhs
        names[f"solver_rel{k}"] = str(rel)
        names[f"solver_rhs{k}"] = rhs if rel == 4 else (q + rhs if "$" in rhs else rhs)
    for n, v in names.items():
        ws.defined_names[n] = DefinedName(n, attr_text=v, hidden=True, localSheetId=idx)

def box(ws, rng, fill=None, b=True):
    for row in ws[rng]:
        for c in row:
            if b: c.border = B
            if fill: c.fill = fill

def sheet(ws, idx, title, var_names, var_values, obj_label, obj_coef, sense_text, typ, rows, neg=1, extra=None, note=None):
    ws["A1"] = title; ws["A1"].font = Font(bold=True, size=13)
    ws["A3"] = "Переменная"; ws["B3"] = "Значение"; ws["C3"] = "Коэф. в целевой"
    for c in "ABC": ws[f"{c}3"].font = bold
    box(ws, "A3:C3", HEAD)
    n = len(var_names)
    for k, (nm, v, cf) in enumerate(zip(var_names, var_values, obj_coef)):
        r = 4 + k
        ws[f"A{r}"] = nm; ws[f"B{r}"] = v; ws[f"C{r}"] = cf
        ws[f"B{r}"].fill = VAR
    box(ws, f"A4:C{3+n}")
    vr = f"$B$4:$B${3+n}"; cr = f"$C$4:$C${3+n}"
    o = 5 + n
    ws[f"A{o}"] = obj_label; ws[f"B{o}"] = f"=SUMPRODUCT({vr},{cr})"; ws[f"C{o}"] = sense_text
    ws[f"A{o}"].font = bold; ws[f"B{o}"].font = bold; box(ws, f"A{o}:C{o}", OBJ)
    h = o + 2
    hdr = ["Ограничение", "Левая часть"] + [f"коэф. {v}" for v in var_names] + ["Знак", "Правая часть"]
    for j, t in enumerate(hdr):
        cell = ws.cell(row=h, column=1 + j, value=t); cell.font = bold
    last_col = openpyxl.utils.get_column_letter(len(hdr))
    box(ws, f"A{h}:{last_col}{h}", HEAD)
    cons = []
    for i, (label, coefs, sign, rhs, rel) in enumerate(rows):
        r = h + 1 + i
        ws.cell(row=r, column=1, value=label)
        c1 = openpyxl.utils.get_column_letter(3); c2 = openpyxl.utils.get_column_letter(2 + n)
        ws.cell(row=r, column=2, value=f"=SUMPRODUCT(TRANSPOSE({vr}),{c1}{r}:{c2}{r})")
        for j, a in enumerate(coefs):
            ws.cell(row=r, column=3 + j, value=a)
        ws.cell(row=r, column=3 + n, value=sign)
        ws.cell(row=r, column=4 + n, value=rhs)
        rhs_col = openpyxl.utils.get_column_letter(4 + n)
        cons.append((f"$B${r}", rel, f"${rhs_col}${r}"))
    box(ws, f"A{h+1}:{last_col}{h+len(rows)}")
    # SUMPRODUCT(TRANSPOSE(..)) требует массивной формулы — используем явную сумму произведений
    for i in range(len(rows)):
        r = h + 1 + i
        terms = [f"{openpyxl.utils.get_column_letter(3+j)}{r}*$B${4+j}" for j in range(n)]
        ws.cell(row=r, column=2, value="=" + "+".join(terms))
    if extra:
        cons += extra(vr)
    solver_model(ws, idx, f"$B${o}", vr, typ, cons, neg)
    ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 13
    for j in range(3, 5 + n):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 11
    ws.column_dimensions["C"].width = 15
    r = h + len(rows) + 2
    lines = note or []
    for k, t in enumerate(lines):
        ws.cell(row=r + k, column=1, value=t).font = Font(italic=True, color="555555")
    return o

rows = [
    ("1: x1 + 2x2 + x4 <= 10", [1, 2, 0, 1], "<=", 10, 1),
    ("2: x1 + x3 + x4 = 7",    [1, 0, 1, 1], "=",  7,  2),
    ("3: x2 + 2x3 >= 5",       [0, 1, 2, 0], ">=", 5,  3),
]
how = ["Данные → Поиск решения (Solver): модель уже сохранена в листе —",
       "целевая ячейка, изменяемые ячейки и ограничения подставятся сами.",
       "Метод решения: «Поиск решения лин. задач симплекс-методом» (Simplex LP) → Найти решение."]

ws = wb.active; ws.title = "ЗЛП"
sheet(ws, 0, "ЛР1. Вариант 6: Z = 2x1 + x2 + x3 + 3x4 → min", ["x1", "x2", "x3", "x4"], [0, 0, 7, 0],
      "Z (целевая функция)", [2, 1, 1, 3], "→ min", 2, rows,
      note=how + ["Ответ: x* = (0, 0, 7, 0), Z(x*) = 7."])

ws2 = wb.create_sheet("Двойственная")
drows = [
    ("при x1: -y1 + y2 <= 2",   [-1, 1, 0], "<=", 2, 1),
    ("при x2: -2y1 + y3 <= 1",  [-2, 0, 1], "<=", 1, 1),
    ("при x3: y2 + 2y3 <= 1",   [0, 1, 2],  "<=", 1, 1),
    ("при x4: -y1 + y2 <= 3",   [-1, 1, 0], "<=", 3, 1),
]
sheet(ws2, 1, "Бонус: двойственная задача  W = -10y1 + 7y2 + 5y3 → max", ["y1", "y2", "y3"], [0, 1, 0],
      "W (целевая функция)", [-10, 7, 5], "→ max", 1, drows, neg=2,
      extra=lambda vr: [("$B$4", 3, "0"), ("$B$6", 3, "0")],
      note=["Ограничение 1 прямой задачи (<=) умножено на -1 → y1 >= 0; ограничение 2 (=) → y2 любого знака;",
            "поэтому флажок «Сделать переменные без ограничений неотрицательными» снят, y1 >= 0 и y3 >= 0 заданы явно.",
            "Ответ: y* = (0, 1, 0), W(y*) = 7 = Z(x*) — теорема двойственности."])

ws3 = wb.create_sheet("Целочисленная")
sheet(ws3, 2, "Бонус: целочисленная задача (x1..x4 — целые)", ["x1", "x2", "x3", "x4"], [0, 0, 7, 0],
      "Z (целевая функция)", [2, 1, 1, 3], "→ min", 2, rows,
      extra=lambda vr: [(vr, 4, "\"integer\"")],
      note=["Добавлено ограничение $B$4:$B$7 = целое (int).",
            "Оптимум ЛП-задачи уже целочисленный, поэтому ответ тот же: x* = (0, 0, 7, 0), Z = 7."])

wb.save(__import__("pathlib").Path(__file__).resolve().parent.parent / "lr1" / "excel_solution.xlsx")
print("saved")
